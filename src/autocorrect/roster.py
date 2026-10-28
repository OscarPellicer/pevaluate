"""Class lists (rosters): the real students of a course, read from the LMS export, and the
matching of the students found in exams or submissions to them.

Supported exports: a PoliformaT "Descargar todo" (zip, folder or its ``grades.xls``/``grades.csv``),
a Moodle grading worksheet or gradebook export (CSV/XLSX), or any CSV/XLSX with an ID column and
name column(s). The matching (ID first, then the most similar name, one-to-one) is the one used
by ``poliformat-zip``.
"""
import csv
import tempfile
from pathlib import Path
from typing import Optional

from .names import name_similarity, normalize_student_id

ID_COLUMNS = ("ID", "Identificador", "Número de ID", "Número de identificación", "Número d'identificació", "ID number", "DNI", "NIF", "NIA", "student_id", "id",
              "Display ID", "Usuario", "Username")
NAME_COLUMNS = ("Nombre completo", "Nombre alumno", "Nom complet", "Full name", "student_name", "name", "Nombre y apellidos", "Alumno", "Student")
FIRST_NAME_COLUMNS = ("Nombre", "Nom", "First name", "First Name")
SURNAME_COLUMNS = ("Apellido(s)", "Apellidos", "Apellido", "Cognoms", "Cognom", "Surname", "Last name", "Last Name")
EMAIL_COLUMNS = ("Dirección de correo", "Correo electrónico", "Email address", "Email", "email")


def _column(header: list, *names) -> Optional[int]:
    lowered = [str(h).strip().lower() for h in header]
    for name in names:
        if name.lower() in lowered:
            return lowered.index(name.lower())
    return None


def _rows_of(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        sheet = xlrd.open_workbook(str(path)).sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                    value = int(value)
                row.append("" if cell.ctype == xlrd.XL_CELL_EMPTY else str(value))
            rows.append(row)
        return rows
    if suffix == ".xlsx":
        import openpyxl

        sheet = openpyxl.load_workbook(str(path), read_only=True, data_only=True).worksheets[0]
        return [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    return list(csv.reader(text.splitlines(), dialect))


def _find_table(path: Path) -> Path:
    """The grades file inside a PoliformaT download (zip or folder), or ``path`` itself."""
    if path.is_dir():
        found = sorted(path.rglob("grades.xls")) + sorted(path.rglob("grades.csv"))
        if not found:
            raise ValueError(f"{path}: no grades.xls or grades.csv inside.")
        return found[0]
    return path


def load_roster(path: str, id_column: Optional[str] = None, name_column: Optional[str] = None) -> list[dict]:
    """Students of a class list as ``{id, name, email}``, ``name`` as 'Name Surname'. ``id_column`` and
    ``name_column`` name the columns when they are not among the usual ones (as in a course record)."""
    source = Path(path)
    columns = {"id_name": id_column, "name_name": name_column}
    if source.suffix.lower() == ".zip":
        import zipfile

        with tempfile.TemporaryDirectory() as tmp:
            with zipfile.ZipFile(source) as archive:
                members = [m for m in archive.namelist() if m.endswith(("grades.xls", "grades.csv"))]
                if not members:
                    raise ValueError(f"{path}: no grades.xls or grades.csv inside.")
                archive.extract(members[0], tmp)
            return _load_table(Path(tmp) / members[0], **columns)
    return _load_table(_find_table(source), **columns)


def _load_table(path: Path, id_name: Optional[str] = None, name_name: Optional[str] = None) -> list[dict]:
    rows = _rows_of(path)
    for index, header in enumerate(rows):
        id_col = _column(header, *([id_name] if id_name else []), *ID_COLUMNS)
        name_col = _column(header, *([name_name] if name_name else []), *NAME_COLUMNS)
        first_col = _column(header, *FIRST_NAME_COLUMNS)
        surname_col = _column(header, *SURNAME_COLUMNS)
        if id_col is None and name_col is None and first_col is None:
            continue
        email_col = _column(header, *EMAIL_COLUMNS)
        students = []
        for row in rows[index + 1:]:
            cell = lambda col: str(row[col]).strip() if col is not None and col < len(row) else ""
            if name_col is not None and cell(name_col):
                name = cell(name_col)
                if "," in name:  # "Surname, Name"
                    surname, _, first = name.partition(",")
                    name = f"{first.strip()} {surname.strip()}"
            else:
                name = f"{cell(first_col)} {cell(surname_col)}".strip()
            student_id = cell(id_col)
            if not name and not student_id:
                continue
            students.append({"id": student_id, "name": name, "email": cell(email_col)})
        return students
    raise ValueError(f"{path}: no header with an ID or name column.")


def match_people(people: list[dict], roster: list[dict], name_threshold: float = 70.0) -> dict:
    """``{index in people: (index in roster, method, score)}``, one-to-one: first by ID (one
    contains the other, e.g. the DNI digits inside a DNI read with its letter), then by the most
    similar name. ``people`` and ``roster`` items have ``id`` and ``name``."""
    pairs = []
    for i, person in enumerate(people):
        our_id = normalize_student_id(person.get("id"))
        for j, student in enumerate(roster):
            their_id = normalize_student_id(student.get("id"))
            if len(our_id) >= 5 and len(their_id) >= 5 and (our_id in their_id or their_id in our_id):
                pairs.append((200.0, i, j, "id"))
            elif person.get("name") and student.get("name"):
                score = name_similarity(person["name"], student["name"])
                if score >= name_threshold:
                    pairs.append((score, i, j, "name"))
    matched, used_people, used_roster = {}, set(), set()
    for score, i, j, method in sorted(pairs, key=lambda p: -p[0]):
        if i in used_people or j in used_roster:
            continue
        matched[i] = (j, method, min(score, 100.0))
        used_people.add(i)
        used_roster.add(j)
    return matched


# --- `pevaluate roster`: one class-list CSV for the course folder ----------------------------------

DESCRIPTION = """Write the class list of an LMS export as one CSV (id, name, email) to keep in the course
folder: pexams correct takes it as --input-csv (--id-column id --name-column name), so the students
are matched to the list while correcting, and the review app and the uploads use it too.

    pevaluate roster "Tarea_20260101120000.zip" --output lista_clase.csv
"""


def build_parser(parser=None):
    import argparse

    parser = parser or argparse.ArgumentParser(prog="pevaluate roster", description=DESCRIPTION,
                                               formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("export", help="PoliformaT 'Descargar todo' (zip, folder, grades.xls/csv), Moodle export, or any CSV/XLSX with ID and name columns.")
    parser.add_argument("--output", required=True, help="The CSV to write (columns id, name, email).")
    parser.add_argument("--id-column", help="ID column of the export, if not one of the usual ones.")
    parser.add_argument("--name-column", help="Name column of the export, if not one of the usual ones.")
    return parser


def run(args):
    students = load_roster(args.export, args.id_column, args.name_column)
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    with open(args.output, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["id", "name", "email"])
        writer.writeheader()
        writer.writerows({k: s.get(k, "") for k in ("id", "name", "email")} for s in students)
    print(f"{len(students)} students written to {args.output}")
