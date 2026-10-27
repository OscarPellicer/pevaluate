"""Word (.docx) and PowerPoint (.pptx) to Markdown, and the media images inside their zip."""
import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from .images import (
    MarkdownExtraction,
    _collect_markdown_asset_images,
    _guess_mime,
    _image_from_bytes,
    _write_image_asset,
)


def _paragraph_to_markdown(para, image_output_dir, image_reference_dir):
    parts = []
    for run in para.runs:
        parts.append(run.text)
        blips = run.element.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
        for blip in blips:
            rel_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if not rel_id:
                continue
            part = para.part.related_parts.get(rel_id)
            if not part:
                continue
            image = _image_from_bytes(str(part.partname), part.blob, mime_type=getattr(part, "content_type", None))
            if image_output_dir:
                parts.append(f"\n\n![{Path(str(part.partname)).name}]({_write_image_asset(image, image_output_dir, image_reference_dir)})\n\n")
            else:
                parts.append("\n\n[image omitted]\n\n")
    text = "".join(parts).strip()
    style_name = getattr(getattr(para, "style", None), "name", "") or ""
    match = re.match(r"Heading\s+([1-6])", style_name, flags=re.I)
    if match and text:
        return f"{'#' * int(match.group(1))} {text}"
    return text


def _docx_iter_blocks(doc):
    try:
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return []
    parent_elm = doc.element.body
    for child in parent_elm.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _docx_table_to_markdown(table):
    rows = [[cell.text.strip().replace("\n", "<br>") for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _read_docx_markdown(file_path, image_output_dir, image_reference_dir):
    try:
        import docx
    except ImportError:
        return MarkdownExtraction("Error: python-docx not installed. Cannot read .docx files.", [], [])

    doc = docx.Document(file_path)
    blocks = []
    for block in _docx_iter_blocks(doc):
        if hasattr(block, "runs"):
            text = _paragraph_to_markdown(block, image_output_dir, image_reference_dir)
        else:
            text = _docx_table_to_markdown(block)
        if text:
            blocks.append(text)
    markdown = "\n\n".join(blocks)
    images, image_paths = _collect_markdown_asset_images(markdown, image_output_dir, image_reference_dir)
    return MarkdownExtraction(markdown, images, [], image_paths)


def _pptx_shape_text(shape):
    text_parts = []
    if hasattr(shape, "text"):
        shape_text = str(shape.text or "").strip()
        if shape_text:
            text_parts.append(shape_text)
    if getattr(shape, "has_table", False):
        rows = []
        for row in shape.table.rows:
            rows.append([cell.text_frame.text.strip().replace("\n", "<br>") for cell in row.cells])
        if rows:
            width = max(len(row) for row in rows)
            rows = [row + [""] * (width - len(row)) for row in rows]
            lines = [
                "| " + " | ".join(rows[0]) + " |",
                "| " + " | ".join("---" for _ in rows[0]) + " |",
            ]
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")
            text_parts.append("\n".join(lines))
    return "\n\n".join(part for part in text_parts if part)


def _read_pptx_markdown(file_path, image_output_dir, image_reference_dir):
    try:
        from pptx import Presentation
    except ImportError:
        Presentation = None

    if Presentation is not None:
        try:
            prs = Presentation(file_path)
            slides = []
            for slide_no, slide in enumerate(prs.slides, start=1):
                parts = [f"## Slide {slide_no}"]
                for shape in slide.shapes:
                    text = _pptx_shape_text(shape)
                    if text:
                        parts.append(text)
                    if getattr(shape, "shape_type", None) is not None and hasattr(shape, "image"):
                        try:
                            image_blob = shape.image.blob
                            mime_type = shape.image.content_type
                            image = _image_from_bytes(
                                f"{file_path}:slide-{slide_no}:{getattr(shape, 'name', 'image')}",
                                image_blob,
                                mime_type=mime_type,
                            )
                            if image_output_dir:
                                parts.append(f"![{getattr(shape, 'name', 'image')}]({_write_image_asset(image, image_output_dir, image_reference_dir)})")
                            else:
                                parts.append("[image omitted]")
                        except Exception:
                            pass
                slides.append("\n\n".join(part for part in parts if part))
            markdown = "\n\n".join(slides)
            images, image_paths = _collect_markdown_asset_images(markdown, image_output_dir, image_reference_dir)
            return MarkdownExtraction(markdown, images, [], image_paths)
        except Exception:
            pass

    ns = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    slides = []
    images = []
    with zipfile.ZipFile(file_path) as zf:
        slide_names = sorted(
            (name for name in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", name)),
            key=lambda name: int(re.search(r"slide(\d+)\.xml$", name).group(1)),
        )
        for slide_no, slide_name in enumerate(slide_names, start=1):
            rels_name = f"ppt/slides/_rels/{Path(slide_name).name}.rels"
            rel_targets = {}
            if rels_name in zf.namelist():
                rel_root = ET.fromstring(zf.read(rels_name))
                for rel in rel_root:
                    rel_id = rel.attrib.get("Id")
                    target = rel.attrib.get("Target", "")
                    rel_targets[rel_id] = target

            root = ET.fromstring(zf.read(slide_name))
            parts = [f"## Slide {slide_no}"]
            sp_tree = root.find(".//p:cSld/p:spTree", ns)
            for element in list(sp_tree) if sp_tree is not None else []:
                texts = [node.text for node in element.findall(".//a:t", ns) if node.text]
                if texts:
                    parts.append("\n".join(texts))
                for blip in element.findall(".//a:blip", ns):
                    rel_id = blip.attrib.get(f"{{{ns['r']}}}embed")
                    target = rel_targets.get(rel_id)
                    if not target:
                        continue
                    if target.startswith("../"):
                        media_name = "ppt/" + target[3:]
                    else:
                        media_name = "ppt/slides/" + target
                    if not media_name.startswith("ppt/"):
                        media_name = "ppt/" + target.lstrip("../")
                    if media_name not in zf.namelist():
                        media_name = "ppt/media/" + Path(target).name
                    if media_name not in zf.namelist():
                        continue
                    data = zf.read(media_name)
                    image = _image_from_bytes(f"{file_path}:{media_name}", data, mime_type=_guess_mime(media_name, data))
                    images.append(image)
                    if image_output_dir:
                        parts.append(f"![{Path(media_name).name}]({_write_image_asset(image, image_output_dir, image_reference_dir)})")
                    else:
                        parts.append("[image omitted]")
            slides.append("\n\n".join(part for part in parts if part))
    markdown = "\n\n".join(slides)
    image_paths = []
    if image_output_dir:
        images, image_paths = _collect_markdown_asset_images(markdown, image_output_dir, image_reference_dir)
    return MarkdownExtraction(markdown, images, [], image_paths)


def _read_zip_media_images(file_path, media_prefixes):
    images = []
    with zipfile.ZipFile(file_path) as zf:
        for member in zf.infolist():
            member_name = member.filename.replace("\\", "/")
            if member.is_dir() or not member_name.lower().startswith(media_prefixes):
                continue
            mime_type = _guess_mime(member_name)
            if not mime_type.startswith("image/"):
                continue
            data = zf.read(member)
            images.append(_image_from_bytes(f"{file_path}:{member_name}", data, mime_type=mime_type))
    return images
